"""Small python-docx helpers shared by the report and worksheet builders."""
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
ACCENT = RGBColor(0x1C, 0x4E, 0x80)
HEAD_FILL = "E6EDF5"
CODE_FILL = "F4F4F2"
BORDER = "B8BEC6"
BODY_FONT = "Calibri"
MONO_FONT = "Consolas"


def new_document(body_pt=10.5):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.9)

    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    st.font.size = Pt(body_pt)
    st.font.color.rgb = INK
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.12

    for name, size, before, after in (("Heading 1", 16, 16, 6), ("Heading 2", 12.5, 12, 4), ("Heading 3", 11, 9, 3)):
        h = doc.styles[name]
        h.font.name = BODY_FONT
        h.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        h.element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        h.element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = ACCENT if name != "Heading 3" else INK
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        doc.styles[name].font.name = BODY_FONT
        doc.styles[name].paragraph_format.space_after = Pt(2)
    return doc


def _field(run, instr):
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == "1":
            el = OxmlElement("w:t")
            el.text = text
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def add_footer(doc, left_text):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.text = ""
    r = p.add_run(left_text + "    |    Page ")
    r.font.size = Pt(8.5)
    r.font.color.rgb = MUTED
    r2 = p.add_run()
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = MUTED
    _field(r2, "PAGE")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_toc(doc, levels="1-2"):
    from docx.enum.style import WD_STYLE_TYPE
    for name, indent in (("TOC 1", 0), ("TOC 2", 0.5)):      # compact entries so the list stays on one page
        try:
            st = doc.styles[name]
        except KeyError:
            st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["Normal"]
        st.font.size = Pt(10)
        st.paragraph_format.space_after = Pt(2.5)
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.line_spacing = 1.0
        st.paragraph_format.left_indent = Cm(indent)
    p = doc.add_paragraph()
    _field(p.add_run(), f'TOC \\o "{levels}" \\h \\z \\u')
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def rich(p, text, size=None, color=None, mono_size=None):
    """Add text to a paragraph; **bold** and `code` markers are honoured."""
    import re
    for part in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = MONO_FONT
            r._r.rPr.rFonts.set(qn("w:eastAsia"), MONO_FONT)
            r.font.size = Pt(mono_size or ((size or 10.5) - 1))

        else:
            r = p.add_run(part)
        if size and not (part.startswith("`") and part.endswith("`")):
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
    return p


def para(doc, text, size=None, align=None, after=None, before=None, color=None, keep=False):
    p = doc.add_paragraph()
    rich(p, text, size=size, color=color)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def bullets(doc, items, style="List Bullet", size=None):
    for it in items:
        p = doc.add_paragraph(style=style)
        rich(p, it, size=size)
        p.paragraph_format.space_after = Pt(2)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(6)


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _cell_margins(cell, top=40, bottom=40, left=70, right=70):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def _borders(table, color=BORDER, size=4):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def table(doc, headers, rows, widths, size=8.5, align=None, bold_first_col=False, total_row=False, mono_cols=(),
          keep_together=False):
    """widths in cm; align is a string per column made of l / r / c."""
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _borders(t)
    align = align or "l" * len(headers)
    amap = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER}

    def fill(cell, text, i, header=False, bold=False):
        cell.width = Cm(widths[i])
        _cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.alignment = amap[align[i]] if not header or align[i] != "l" else WD_ALIGN_PARAGRAPH.LEFT
        if header:
            r = p.add_run(str(text))
            r.bold = True
            r.font.size = Pt(size)
            _shade(cell, HEAD_FILL)
        else:
            if i in mono_cols:
                r = p.add_run(str(text))
                r.font.name = MONO_FONT
                r._r.rPr.rFonts.set(qn("w:eastAsia"), MONO_FONT)
                r.font.size = Pt(size - 0.5)
                r.bold = bold
            else:
                rich(p, str(text), size=size, mono_size=size - 0.5)
                if bold:
                    for r in p.runs:
                        r.bold = True

    for i, h in enumerate(headers):
        fill(t.rows[0].cells[i], h, i, header=True)
    _repeat_header(t.rows[0])
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        last = total_row and ri == len(rows) - 1
        for i, v in enumerate(row):
            fill(cells[i], v, i, bold=last or (bold_first_col and i == 0))
        if last:
            for c in cells:
                _shade(c, "F2F2F0")
        _no_split(t.rows[-1])
    if keep_together:                      # chain every row to the next so the table is not split over pages
        for row in t.rows[:-1]:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    spacer.paragraph_format.line_spacing = 0.6
    return t


def caption(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(9)
    r.italic = True
    r.font.color.rgb = MUTED
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    return p


def code(doc, text, size=8):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(t, color="D9D9D6")
    cell = t.rows[0].cells[0]
    _shade(cell, CODE_FILL)
    _cell_margins(cell, 70, 70, 110, 110)
    lines = text.strip("\n").split("\n")
    p = cell.paragraphs[0]
    for i, line in enumerate(lines):
        if i:
            p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line if line else " ")
        r.font.name = MONO_FONT
        r._r.rPr.rFonts.set(qn("w:eastAsia"), MONO_FONT)
        r.font.size = Pt(size)
    _no_split(t.rows[0])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    spacer.paragraph_format.line_spacing = 0.6


def figure(doc, path, width_cm, cap):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(cap)
    r.font.size = Pt(9)
    r.italic = True
    r.font.color.rgb = MUTED
    c.paragraph_format.space_after = Pt(8)
